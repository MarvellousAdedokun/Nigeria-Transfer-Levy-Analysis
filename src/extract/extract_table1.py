#!/usr/bin/env python3
"""
Extract Table I EMTL cells from FAAC Excel workbooks.

Defaults:
  input:  data/raw/faac/
  output: data/interim/staging_table1.csv
          data/interim/source_register.csv
          docs/extraction-log.md
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

EMTL_RE = re.compile(r"electronic\s+money\s+transfer|emtl", re.I)
TITLE_RE = re.compile(r"\bmonth\s+of\s+([A-Za-z]+\s*,?\s*\d{4})", re.I)
SHARED_RE = re.compile(r"\bshared\s+in\s+([A-Za-z]+\s*,?\s*\d{4})", re.I)
FILENAME_PERIOD_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)[,_\s-]+(20\d{2})",
    re.I,
)
TOTAL_RE = re.compile(r"^\s*total\s*$", re.I)
TABLE_RE = re.compile(r"^\s*Table\s+II\b", re.I)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def raw_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def number_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    if isinstance(value, str):
        s = value.strip().replace(",", "").replace("₦", "")
        if not s or s == "₦":
            return None
        try:
            return Decimal(s)
        except InvalidOperation:
            return None
    return None


def money2(value: Decimal | None) -> Decimal | None:
    if value is None:
        return None
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def iter_source_files(raw_dir: Path):
    """Yield source records for physical files and archive members.

    Workbook bytes are retained only for .xlsx inputs. Other files are hashed
    and registered without keeping their contents in memory.
    """
    for path in sorted(p for p in raw_dir.rglob("*") if p.is_file()):
        relative = path.relative_to(raw_dir).as_posix()
        if path.suffix.lower() == ".zip":
            archive_hash = sha256_file(path)
            archive_size = path.stat().st_size
            # Register the original archive itself as well as its contents.
            yield {
                "kind": "physical_file",
                "source_path": relative,
                "source_file": path.name,
                "source_container": "",
                "archive_sha256": "",
                "member_sha256": archive_hash,
                "size_bytes": archive_size,
                "extension": path.suffix.lower(),
                "bytes": None,
            }
            try:
                with zipfile.ZipFile(path) as zf:
                    for info in zf.infolist():
                        if info.is_dir():
                            continue
                        ext = Path(info.filename).suffix.lower()
                        member_bytes = zf.read(info.filename)
                        yield {
                            "kind": "archive_member",
                            "source_path": f"{relative}::{info.filename}",
                            "source_file": Path(info.filename).name,
                            "source_container": relative,
                            "archive_sha256": archive_hash,
                            "member_sha256": sha256_bytes(member_bytes),
                            "size_bytes": len(member_bytes),
                            "extension": ext,
                            "bytes": member_bytes if ext == ".xlsx" else None,
                        }
            except zipfile.BadZipFile:
                yield {
                    "kind": "physical_file",
                    "source_path": relative,
                    "source_file": path.name,
                    "source_container": "",
                    "archive_sha256": "",
                    "member_sha256": archive_hash,
                    "size_bytes": archive_size,
                    "extension": path.suffix.lower(),
                    "bytes": None,
                }
        else:
            ext = path.suffix.lower()
            digest = sha256_file(path)
            size = path.stat().st_size
            data = path.read_bytes() if ext == ".xlsx" else None
            yield {
                "kind": "physical_file",
                "source_path": relative,
                "source_file": path.name,
                "source_container": "",
                "archive_sha256": "",
                "member_sha256": digest,
                "size_bytes": size,
                "extension": ext,
                "bytes": data,
            }


def make_source_register(records: list[dict[str, Any]], output: Path) -> None:
    fields = [
        "source_path", "source_file", "kind", "source_container",
        "extension", "size_bytes", "sha256", "archive_sha256", "notes",
    ]
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for r in records:
            if r["kind"] == "physical_file":
                digest = r["member_sha256"]
            else:
                digest = r["member_sha256"]
            writer.writerow({
                "source_path": r["source_path"],
                "source_file": r["source_file"],
                "kind": r["kind"],
                "source_container": r["source_container"],
                "extension": r["extension"],
                "size_bytes": r["size_bytes"],
                "sha256": digest,
                "archive_sha256": r["archive_sha256"],
                "notes": "ZIP member checksum is for the member bytes; archive checksum is recorded separately."
                if r["kind"] == "archive_member" else "",
            })


def find_sheet(wb):
    for name in wb.sheetnames:
        low = name.strip().lower()
        if low.startswith("sum") and "fg" in low:
            return wb[name]
    return None


def find_table_i_start(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15)):
        for c in row:
            if isinstance(c.value, str) and re.fullmatch(r"\s*Table\s+I\s*", c.value, re.I):
                return c.row
    return 1


def find_title(ws, table_start: int) -> tuple[str, str]:
    for row in ws.iter_rows(min_row=table_start, max_row=min(ws.max_row, table_start + 5)):
        for c in row:
            if isinstance(c.value, str) and "Summary of Gross Revenue Allocation" in c.value:
                return c.coordinate, c.value.strip()
    # fallback: preserve the likely title cell if present
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
        for c in row:
            if isinstance(c.value, str) and re.search(r"\bmonth\s+of\b", c.value, re.I):
                return c.coordinate, c.value.strip()
    return "", ""


def find_table_i_end(ws, table_start: int) -> int:
    """Return the row before Table II, or the last used row if Table II is absent."""
    for r in range(table_start + 1, ws.max_row + 1):
        if TABLE_RE.search(raw_text(ws.cell(r, 1).value)):
            return r - 1
    return ws.max_row


def find_header_row(ws, table_start: int, title_cell: str) -> tuple[int | None, int | None, str]:
    # Prefer a row with both Beneficiaries and Total; search only within Table I.
    start = table_start + 1
    end = find_table_i_end(ws, table_start)
    for r in range(start, end + 1):
        vals = [raw_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        has_beneficiaries = any(v.lower() == "beneficiaries" for v in vals)
        has_total = any(v.lower() == "total" for v in vals)
        if has_beneficiaries and has_total:
            emtl_cell = None
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and EMTL_RE.search(v):
                    emtl_cell = c
                    break
            return r, emtl_cell, raw_text(ws.cell(r, emtl_cell).value) if emtl_cell else ""
    return None, None, ""


def find_total_row(ws, header_row: int | None) -> int | None:
    if header_row is None:
        return None
    for r in range(header_row + 1, ws.max_row + 1):
        # Table I TOTAL label is in column B. Stop if the next table begins first.
        a = raw_text(ws.cell(r, 1).value)
        b = raw_text(ws.cell(r, 2).value)
        if TABLE_RE.search(a):
            break
        if TOTAL_RE.fullmatch(b):
            return r
    return None


def filename_period(filename: str) -> str:
    m = FILENAME_PERIOD_RE.search(filename.replace(".", " "))
    if not m:
        return ""
    return f"{m.group(1)}, {m.group(2)}"


def get_periods(title: str, source_file: str) -> tuple[str, str, str]:
    revenue = ""
    allocation = ""
    allocation_source = ""
    m = TITLE_RE.search(title or "")
    if m:
        revenue = m.group(1).strip()
    m = SHARED_RE.search(title or "")
    if m:
        allocation = m.group(1).strip()
        allocation_source = "title"
    else:
        allocation = filename_period(source_file)
        allocation_source = "filename_fallback" if allocation else "not_stated"
    return revenue, allocation, allocation_source


def extract_workbook(source: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    data = source["bytes"]
    wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    ws = find_sheet(wb)
    warnings: list[str] = []
    base = {
        "source_file": source["source_file"],
        "source_path": source["source_path"],
        "source_container": source["source_container"],
        "sha256": source["member_sha256"],
        "archive_sha256": source["archive_sha256"],
    }

    if ws is None:
        warning = "No sheet matching 'Sum & FG' was found."
        row = {
            **base, "sheet": "", "title_cell": "", "title_raw": "",
            "revenue_period_raw": "", "allocation_period_raw": "",
            "allocation_period_source": "", "header_row": "",
            "header_cell": "", "header_label_raw": "",
            "row_label_cell": "", "row_label_raw": "", "cell": "",
            "value_raw": "", "row_type": "extraction_error",
            "emtl_found": False, "status": "sheet_not_found",
            "total_label_cell": "", "emtl_total_cell": "", "emtl_total_value_raw": "",
            "overall_total_cell": "", "overall_total_value_raw": "", "detail_sum_raw": "",
            "sum_matches_total": "", "warning": warning,
        }
        return [row], {"source_file": source["source_file"], "source_path": source["source_path"],
                       "status": "sheet_not_found", "rows_extracted": 1, "warning": warning}

    table_start = find_table_i_start(ws)
    title_cell, title = find_title(ws, table_start)
    revenue_period, allocation_period, allocation_source = get_periods(title, source["source_file"])
    header_row, emtl_col, header_label = find_header_row(ws, table_start, title_cell)
    total_row = find_total_row(ws, header_row)

    if not title:
        warnings.append("Table I title text was not found in the expected area.")
    if header_row is None:
        warnings.append("Table I heading row was not identified.")
    if total_row is None:
        warnings.append("Table I TOTAL row was not identified.")

    # Find the overall Total column from the Table I header row.
    total_col = None
    if header_row:
        for c in range(1, ws.max_column + 1):
            if raw_text(ws.cell(header_row, c).value).lower() == "total":
                total_col = c
                break
    total_label_cell = f"B{total_row}" if total_row else ""
    overall_total_cell = f"{ws.cell(total_row, total_col).coordinate}" if total_row and total_col else ""
    overall_total_value = ws.cell(total_row, total_col).value if total_row and total_col else None
    emtl_total_cell = f"{ws.cell(total_row, emtl_col).coordinate}" if total_row and emtl_col else ""
    emtl_total_value = ws.cell(total_row, emtl_col).value if total_row and emtl_col else None

    rows: list[dict[str, Any]] = []
    status = "found" if emtl_col else "column_not_found"
    if header_row is None:
        status = "header_not_found"
    elif total_row is None:
        status = "total_row_not_found" if emtl_col else "column_not_found_total_row_not_found"

    # Keep the missing column as an explicit audit record; never replace with zero.
    if not emtl_col:
        warnings.append("No EMTL heading found in the Table I header row; this is missing/not reported, not zero.")
        rows.append({
            **base, "sheet": ws.title, "title_cell": title_cell, "title_raw": title,
            "revenue_period_raw": revenue_period, "allocation_period_raw": allocation_period,
            "allocation_period_source": allocation_source, "header_row": header_row or "",
            "header_cell": "", "header_label_raw": "",
            "row_label_cell": "", "row_label_raw": "", "cell": "",
            "value_raw": "", "row_type": "missing_column",
            "emtl_found": False, "status": status,
            "total_label_cell": total_label_cell, "emtl_total_cell": emtl_total_cell,
            "emtl_total_value_raw": emtl_total_value if emtl_total_value is not None else "",
            "overall_total_cell": overall_total_cell,
            "overall_total_value_raw": overall_total_value if overall_total_value is not None else "",
            "detail_sum_raw": "", "sum_matches_total": "",
            "warning": " | ".join(warnings),
        })
    else:
        detail_values: list[Decimal] = []
        # Extract labelled Table I rows through TOTAL, inclusive.
        stop_row = total_row if total_row else find_table_i_end(ws, table_start)
        for r in range((header_row or table_start) + 1, stop_row + 1):
            label_value = ws.cell(r, 2).value
            emtl_value = ws.cell(r, emtl_col).value
            label = raw_text(label_value)
            numeric = number_decimal(emtl_value)
            is_total = bool(TOTAL_RE.fullmatch(label))
            # Skip currency/unit row and blank spacer rows, but retain rows with labels or numeric values.
            if not label and numeric is None:
                continue
            row_type = "total" if is_total else "detail"
            if not is_total and numeric is not None:
                detail_values.append(numeric)
            rows.append({
                **base, "sheet": ws.title, "title_cell": title_cell, "title_raw": title,
                "revenue_period_raw": revenue_period, "allocation_period_raw": allocation_period,
                "allocation_period_source": allocation_source, "header_row": header_row or "",
                "header_cell": ws.cell(header_row, emtl_col).coordinate,
                "header_label_raw": header_label,
                "row_label_cell": ws.cell(r, 2).coordinate if label_value is not None else "",
                "row_label_raw": label, "cell": ws.cell(r, emtl_col).coordinate,
                "value_raw": emtl_value if emtl_value is not None else "",
                "row_type": row_type, "emtl_found": True, "status": status,
                "total_label_cell": total_label_cell, "emtl_total_cell": emtl_total_cell,
                "emtl_total_value_raw": emtl_total_value if emtl_total_value is not None else "",
                "overall_total_cell": overall_total_cell,
                "overall_total_value_raw": overall_total_value if overall_total_value is not None else "",
                "detail_sum_raw": "", "sum_matches_total": "", "warning": "",
            })

        detail_sum = sum(detail_values, Decimal("0"))
        matches = ""
        if total_row and emtl_col:
            total_decimal = number_decimal(ws.cell(total_row, emtl_col).value)
            matches = (money2(detail_sum) == money2(total_decimal)) if total_decimal is not None else False
            if matches is False:
                warnings.append(
                    f"EMTL detail rows sum to {detail_sum}; Table I EMTL TOTAL is {ws.cell(total_row, emtl_col).value}."
                )
        for row in rows:
            row["detail_sum_raw"] = str(detail_sum)
            row["sum_matches_total"] = matches
            row["warning"] = " | ".join(warnings)

    log = {
        "source_file": source["source_file"],
        "source_path": source["source_path"],
        "sheet": ws.title,
        "title_cell": title_cell,
        "title_raw": title,
        "revenue_period_raw": revenue_period,
        "allocation_period_raw": allocation_period,
        "allocation_period_source": allocation_source,
        "header_row": header_row or "",
        "header_cell": ws.cell(header_row, emtl_col).coordinate if header_row and emtl_col else "",
        "header_label_raw": header_label,
        "total_label_cell": total_label_cell,
        "emtl_total_cell": emtl_total_cell,
        "emtl_total_value_raw": emtl_total_value if emtl_total_value is not None else "",
        "overall_total_cell": overall_total_cell,
        "overall_total_value_raw": overall_total_value if overall_total_value is not None else "",
        "status": status,
        "rows_extracted": len(rows),
        "sum_matches_total": (
            rows[-1].get("sum_matches_total", "")
            if rows and rows[-1].get("row_type") == "total" else ""
        ),
        "manual_fix": "No",
        "warning": " | ".join(warnings),
    }
    return rows, log


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", default="data/raw/faac", help="Folder with original FAAC files and ZIP archives.")
    parser.add_argument("--staging-out", default="data/interim/staging_table1.csv")
    parser.add_argument("--register-out", default="data/interim/source_register.csv")
    parser.add_argument("--log-out", default="docs/extraction-log.md")
    args = parser.parse_args()

    raw_dir = Path(args.raw_dir)
    if not raw_dir.exists():
        raise SystemExit(f"Raw directory does not exist: {raw_dir}")

    source_records = list(iter_source_files(raw_dir))
    register_out = Path(args.register_out)
    make_source_register(source_records, register_out)

    workbook_sources = [r for r in source_records if r["extension"] == ".xlsx"]
    # De-duplicate identical workbook bytes if an extracted copy and ZIP member both exist.
    unique_sources = []
    seen_hashes = {}
    duplicate_sources = []
    for source in workbook_sources:
        digest = source["member_sha256"]
        if digest in seen_hashes:
            duplicate_sources.append({
                "duplicate_path": source["source_path"],
                "processed_path": seen_hashes[digest],
                "sha256": digest,
            })
            continue
        seen_hashes[digest] = source["source_path"]
        unique_sources.append(source)

    staging_rows = []
    log_rows = []
    for source in unique_sources:
        try:
            rows, log_row = extract_workbook(source)
        except Exception as exc:
            warning = f"Extraction failed: {type(exc).__name__}: {exc}"
            rows = [{
                "source_file": source["source_file"], "source_path": source["source_path"],
                "source_container": source["source_container"], "sha256": source["member_sha256"],
                "archive_sha256": source["archive_sha256"], "sheet": "", "title_cell": "",
                "title_raw": "", "revenue_period_raw": "", "allocation_period_raw": "",
                "allocation_period_source": "", "header_row": "", "header_cell": "",
                "header_label_raw": "", "row_label_cell": "", "row_label_raw": "",
                "cell": "", "value_raw": "", "row_type": "extraction_error",
                "emtl_found": False, "status": "extraction_error", "total_label_cell": "",
                "emtl_total_cell": "", "emtl_total_value_raw": "",
                "overall_total_cell": "", "overall_total_value_raw": "", "detail_sum_raw": "",
                "sum_matches_total": "", "warning": warning,
            }]
            log_row = {"source_file": source["source_file"], "source_path": source["source_path"],
                       "sheet": "", "title_cell": "", "title_raw": "", "header_row": "",
                       "header_cell": "", "header_label_raw": "", "total_label_cell": "",
                       "emtl_total_cell": "", "emtl_total_value_raw": "",
                       "overall_total_cell": "", "overall_total_value_raw": "", "status": "extraction_error",
                       "rows_extracted": 1, "warning": warning}
        staging_rows.extend(rows)
        log_rows.append(log_row)

    staging_out = Path(args.staging_out)
    staging_out.parent.mkdir(parents=True, exist_ok=True)
    stage_fields = [
        "source_file", "source_path", "source_container", "sha256", "archive_sha256",
        "sheet", "title_cell", "title_raw", "revenue_period_raw", "allocation_period_raw",
        "allocation_period_source", "header_row", "header_cell", "header_label_raw",
        "row_label_cell", "row_label_raw", "cell", "value_raw", "row_type",
        "emtl_found", "status", "total_label_cell", "emtl_total_cell", "emtl_total_value_raw",
        "overall_total_cell", "overall_total_value_raw", "detail_sum_raw", "sum_matches_total", "warning",
    ]
    with staging_out.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=stage_fields)
        writer.writeheader()
        writer.writerows(staging_rows)

    log_out = Path(args.log_out)
    log_out.parent.mkdir(parents=True, exist_ok=True)
    statuses = Counter(row["status"] for row in log_rows)
    with log_out.open("w", encoding="utf-8") as fh:
        fh.write("# Table I EMTL extraction log\n\n")
        fh.write("This log records what the script found in each Excel workbook. It does not treat a missing EMTL column as zero.\n\n")
        fh.write(f"- Raw directory: `{raw_dir}`\n")
        fh.write(f"- Excel workbooks processed: {len(log_rows)}\n")
        fh.write(f"- Staging rows written: {len(staging_rows)}\n")
        fh.write(f"- Status counts: `{dict(statuses)}`\n\n")
        fh.write("| Source workbook | Sheet | Title cell | Revenue period (raw) | Allocation period (raw) | Heading row | EMTL heading cell | TOTAL label cell | EMTL total cell | Overall total cell | Status | Rows | Sum check | Manual fix | Warning |\n")
        fh.write("|---|---|---|---|---|---:|---|---|---|---|---|---:|---|---|---|\n")
        for r in log_rows:
            vals = [
                r.get("source_path", ""), r.get("sheet", ""), r.get("title_cell", ""),
                r.get("revenue_period_raw", ""), r.get("allocation_period_raw", ""),
                r.get("header_row", ""), r.get("header_cell", ""),
                r.get("total_label_cell", ""), r.get("emtl_total_cell", ""),
                r.get("overall_total_cell", ""), r.get("status", ""), r.get("rows_extracted", ""),
                r.get("sum_matches_total", ""), r.get("manual_fix", "No"),
                r.get("warning", "").replace("|", "/").replace("\\n", " "),
            ]
            fh.write("| " + " | ".join(str(v).replace("|", "/") for v in vals) + " |\n")
        fh.write("\n## Review notes\n\n")
        fh.write("- `found` means the EMTL heading was located in Table I; it does not mean the amount is positive.\n")
        fh.write("- `column_not_found` means the Table I heading row did not contain an EMTL heading. It does not mean EMTL collected or distributed was zero.\n")
        fh.write("- `sum_matches_total` in the staging CSV checks whether the extracted detail cells add to the Table I EMTL TOTAL after rounding to two decimal places.\n")
        fh.write("- Any warning must be reviewed before the staging table is used for analysis.\n")


        fh.write("\\n## Duplicate workbook copies skipped\\n\\n")
        if duplicate_sources:
            fh.write("Identical workbook bytes were found at more than one path. Only one copy was extracted to avoid duplicate observations; all copies remain in the source register.\\n\\n")
            fh.write("| Skipped duplicate path | Processed copy | SHA-256 |\\n|---|---|---|\\n")
            for d in duplicate_sources:
                fh.write(f"| {d['duplicate_path']} | {d['processed_path']} | {d['sha256']} |\\n")
        else:
            fh.write("No duplicate workbook bytes were found.\\n")

    print(f"Workbooks processed: {len(log_rows)}")
    print(f"Staging rows: {len(staging_rows)}")
    print(f"Source register: {register_out}")
    print(f"Staging table: {staging_out}")
    print(f"Extraction log: {log_out}")
    print("Status counts:", dict(statuses))
    print(f"Identical workbook copies skipped: {len(duplicate_sources)}")


if __name__ == "__main__":
    main()
